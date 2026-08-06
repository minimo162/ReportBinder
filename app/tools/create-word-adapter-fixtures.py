#!/usr/bin/env python3
"""Create practical DOCX fixtures for the real-Word source-adapter selfcheck."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLUE = "1F4E78"
PALE_BLUE = "D9EAF7"
PALE_GRAY = "EDF1F5"
GREEN = "E2F0D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge_data in edges.items():
        tag = f"w:{edge_name}"
        tag_obj = tc_borders.find(qn(tag))
        if tag_obj is None:
            tag_obj = OxmlElement(tag)
            tc_borders.append(tag_obj)
        for key, value in edge_data.items():
            tag_obj.set(qn(f"w:{key}"), str(value))


def set_run_font(run, size: float = 10.5, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Yu Gothic"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "游ゴシック")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text: str, size: float = 10.5, bold: bool = False, color: str | None = None):
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, color)
    return run


def configure_section(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)


def decorate_section(section, label: str) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(hp, f"ReportBinder Word Adapter QA  |  {label}", 8.5, False, "5A6B7A")
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp, "CONFIDENTIAL  •  Internal verification fixture", 8, False, "718096")


def add_page_title(doc: Document, eyebrow: str, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_text(p, eyebrow.upper(), 8.5, True, "2878A8")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_text(p, title, 22, True, BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    add_text(p, subtitle, 10.5, False, "526475")


def add_kpi_row(doc: Document, values: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(values))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (label, value, note) in zip(table.rows[0].cells, values):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label + "\n", 8.5, True, "526475")
        add_text(p, value + "\n", 17, True, BLUE)
        add_text(p, note, 8, False, "526475")


def add_prose_page(doc: Document, revision: int) -> None:
    add_page_title(doc, "Executive brief", "部門横断 月次報告", "文章中心の原稿で、見出し・段落・箇条書き・ヘッダー／フッターを検証します。")
    add_kpi_row(doc, [("対象部門", "6", "全提出済み"), ("重要課題", "3", "うち1件更新" if revision > 1 else "前月比 ±0"), ("次回締切", "8/28", "17:00 JST")])
    for heading, body in [
        ("1. エグゼクティブサマリー", "各部門から収集した原稿を統合し、レビュー前の最新版を確定しました。売上は計画線上で推移していますが、調達リードタイムと採用充足率は継続監視が必要です。"),
        ("2. 今月の判断事項", "海外拠点向けの設備投資を二段階承認へ変更し、契約締結前に法務レビューを追加します。更新後の基準は、次回提出分から全案件へ適用します。"),
        ("3. 次回までのアクション", "財務は見通し差異を再確認し、人事は採用計画を更新、総務は提出様式の統一案を配布します。担当者は期日と根拠資料を履歴へ残してください。"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(4)
        add_text(p, heading, 12.5, True, BLUE)
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        add_text(p, body, 10.3)
    for item in ["差異が5%以上の項目はコメントを追記する", "参照ファイル名と更新日時を確認する", "最終出力前にページ順と欠落を確認する"]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        add_text(p, item, 10)


def add_budget_page(doc: Document, revision: int) -> None:
    add_page_title(doc, "Finance", "予算実績サマリー", "複数列・結合見出し・太罫線を含む実務的な表を検証します。")
    headers = ["部門", "予算", "実績", "差異", "進捗", "コメント"]
    rows = [
        ["営業", "¥42.0M", "¥40.8M" if revision == 1 else "¥43.2M", "+¥1.2M" if revision == 1 else "-¥1.2M", "97%" if revision == 1 else "103%", "案件前倒し" if revision > 1 else "計画どおり"],
        ["開発", "¥35.0M", "¥36.4M", "-¥1.4M", "104%", "外部委託増"],
        ["管理", "¥18.0M", "¥17.1M", "+¥0.9M", "95%", "採用時期変更"],
        ["合計", "¥95.0M", "¥94.3M" if revision == 1 else "¥96.7M", "+¥0.7M" if revision == 1 else "-¥1.7M", "99%" if revision == 1 else "102%", "要フォロー" if revision > 1 else "概ね計画内"],
    ]
    table = doc.add_table(rows=2 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    top = table.rows[0].cells
    top[0].merge(top[-1])
    set_cell_shading(top[0], BLUE)
    p = top[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "FY2026 上期 — 部門別執行状況", 11, True, "FFFFFF")
    for i, label in enumerate(headers):
        cell = table.rows[1].cells[i]
        set_cell_shading(cell, PALE_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, 9, True, BLUE)
    for row_index, values in enumerate(rows, start=2):
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            if row_index == len(rows) + 1:
                set_cell_shading(cell, GREEN)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_index in (1, 2, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, value, 9, row_index == len(rows) + 1)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, top={"val": "single", "sz": "5", "color": "AAB7C4"}, bottom={"val": "single", "sz": "5", "color": "AAB7C4"}, left={"val": "single", "sz": "5", "color": "AAB7C4"}, right={"val": "single", "sz": "5", "color": "AAB7C4"})
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    add_text(p, "差異コメント", 12, True, BLUE)
    note = "改訂版では営業実績を更新し、合計が予算超過へ転じたことを明示しています。" if revision > 1 else "実績は概ね計画内です。開発部門の超過要因を次回会議までに精査します。"
    p = doc.add_paragraph()
    add_text(p, note, 10.2)


def add_resume_page(doc: Document) -> None:
    add_page_title(doc, "Human resources", "職務経歴・社内プロフィール", "印刷前提の履歴書型レイアウト、結合セル、固定枠を検証します。")
    table = doc.add_table(rows=7, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    data = [
        ("氏名", "山田 花子", "社員番号", "RB-0264"),
        ("所属", "事業企画部", "役職", "プロジェクトリード"),
        ("連絡先", "hanako.yamada@example.test", "勤務地", "東京"),
        ("専門領域", "業務設計／データ分析", "経験年数", "11年"),
        ("資格", "PMP・簿記2級", "語学", "日本語／英語"),
        ("要約", "部門横断プロジェクトの立ち上げと運用定着を担当。意思決定資料の標準化、業務可視化、進行管理を強みとする。", "", ""),
        ("主要実績", "月次報告の作成工数を35%削減。提出原稿の履歴管理と差分レビューを導入し、確認漏れを半減。", "", ""),
    ]
    for r, values in enumerate(data):
        cells = table.rows[r].cells
        if r >= 5:
            cells[1].merge(cells[3])
            values = (values[0], values[1], "", "")
        for c, value in enumerate(values):
            cell = table.rows[r].cells[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if c in (0, 2) and value:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            add_text(p, value, 9.3, c in (0, 2) and bool(value), BLUE if c in (0, 2) and value else None)
            set_cell_border(cell, top={"val": "single", "sz": "7", "color": "7890A4"}, bottom={"val": "single", "sz": "7", "color": "7890A4"}, left={"val": "single", "sz": "7", "color": "7890A4"}, right={"val": "single", "sz": "7", "color": "7890A4"})
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    add_text(p, "承認欄", 11, True, BLUE)
    approval = doc.add_table(rows=2, cols=3)
    for i, label in enumerate(["本人", "所属長", "人事"]):
        set_cell_shading(approval.rows[0].cells[i], PALE_GRAY)
        approval.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(approval.rows[0].cells[i].paragraphs[0], label, 9, True)
        approval.rows[1].cells[i].height = Cm(1.5)
        set_cell_border(approval.rows[1].cells[i], top={"val": "single", "sz": "6", "color": "7890A4"}, bottom={"val": "single", "sz": "6", "color": "7890A4"}, left={"val": "single", "sz": "6", "color": "7890A4"}, right={"val": "single", "sz": "6", "color": "7890A4"})


def add_landscape_page(doc: Document, revision: int) -> None:
    add_page_title(doc, "Operations", "拠点別リスク・アクション一覧", "横向きページと多列テーブルの保持を検証します。")
    headers = ["ID", "拠点", "リスク", "影響", "確率", "担当", "期限", "状態", "対応方針"]
    rows = [
        ["R-01", "東京", "調達遅延", "高", "中", "佐藤", "8/20", "対応中", "代替ベンダー見積取得"],
        ["R-02", "大阪", "要員不足", "中", "高", "鈴木", "8/25", "要注意", "応援要員を2名確保"],
        ["R-03", "福岡", "設備停止", "高", "低", "田中", "9/02", "監視", "予防保守を前倒し"],
        ["R-04", "札幌", "移行遅延", "中", "中", "高橋", "9/08", "対応中" if revision > 1 else "未着手", "手順レビューを追加"],
        ["R-05", "海外", "契約条件", "高", "中", "伊藤", "9/15", "法務確認", "責任分界点を明文化"],
    ]
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, label in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], BLUE)
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, 8.5, True, "FFFFFF")
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values):
            cell = table.rows[r].cells[c]
            if r % 2 == 0:
                set_cell_shading(cell, "F5F8FA")
            p = cell.paragraphs[0]
            add_text(p, value, 8.2, c == 0)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, bottom={"val": "single", "sz": "4", "color": "B7C4CE"})


def add_approval_page(doc: Document) -> None:
    add_page_title(doc, "Governance", "決裁・配布記録", "縦向きへ戻るセクションと、記入欄を含む帳票形式を検証します。")
    info = doc.add_table(rows=4, cols=4)
    values = [
        ["文書番号", "RB-QA-2026-08", "版", "2.0"],
        ["起案部門", "経営企画部", "起案日", "2026/08/06"],
        ["機密区分", "社内限定", "保存年限", "5年"],
        ["件名", "月次報告パック 最終版承認", "", ""],
    ]
    for r, row in enumerate(values):
        if r == 3:
            info.rows[r].cells[1].merge(info.rows[r].cells[3])
        for c, value in enumerate(row):
            cell = info.rows[r].cells[c]
            if c in (0, 2) and value:
                set_cell_shading(cell, PALE_BLUE)
            add_text(cell.paragraphs[0], value, 9.2, c in (0, 2) and bool(value), BLUE if c in (0, 2) and value else None)
            set_cell_border(cell, top={"val": "single", "sz": "6", "color": "7890A4"}, bottom={"val": "single", "sz": "6", "color": "7890A4"}, left={"val": "single", "sz": "6", "color": "7890A4"}, right={"val": "single", "sz": "6", "color": "7890A4"})
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    add_text(p, "決裁", 12, True, BLUE)
    approval = doc.add_table(rows=3, cols=4)
    for c, label in enumerate(["起案", "部門長", "管理責任者", "最終承認"]):
        set_cell_shading(approval.rows[0].cells[c], PALE_GRAY)
        approval.rows[0].cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(approval.rows[0].cells[c].paragraphs[0], label, 9, True)
        approval.rows[1].cells[c].height = Cm(2.2)
        approval.rows[2].cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(approval.rows[2].cells[c].paragraphs[0], "　　年　月　日", 8.5, False, "526475")
        for row in approval.rows:
            set_cell_border(row.cells[c], top={"val": "single", "sz": "6", "color": "7890A4"}, bottom={"val": "single", "sz": "6", "color": "7890A4"}, left={"val": "single", "sz": "6", "color": "7890A4"}, right={"val": "single", "sz": "6", "color": "7890A4"})
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    add_text(p, "配布先", 11, True, BLUE)
    for item in ["取締役会事務局", "経営企画部", "財務部", "人事部", "各拠点責任者"]:
        p = doc.add_paragraph(style="List Bullet")
        add_text(p, item, 9.8)


def add_appendix_page(doc: Document) -> None:
    add_page_title(doc, "Appendix", "改訂版 追加資料", "ページ追加の検出と履歴比較を確認するための追加ページです。")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    add_text(p, "改訂時に追加された資料です。ページ構成へ新しい項目が同期され、旧版との比較で「追加ページ」として識別されることを検証します。", 11)
    table = doc.add_table(rows=4, cols=3)
    headers = ["確認項目", "担当", "結果"]
    for c, label in enumerate(headers):
        set_cell_shading(table.rows[0].cells[c], BLUE)
        add_text(table.rows[0].cells[c].paragraphs[0], label, 9, True, "FFFFFF")
    for r, row in enumerate([["Word PDF変換", "システム", "合格"], ["更新検出", "システム", "合格"], ["差分履歴", "レビュー担当", "確認待ち"]], start=1):
        for c, value in enumerate(row):
            if r % 2 == 0:
                set_cell_shading(table.rows[r].cells[c], PALE_GRAY)
            add_text(table.rows[r].cells[c].paragraphs[0], value, 9.5, c == 2)


def build_document(path: Path, revision: int) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Yu Gothic"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "游ゴシック")
    styles["Normal"].font.size = Pt(10.5)

    section = doc.sections[0]
    configure_section(section)
    decorate_section(section, "01 / Executive brief")
    add_prose_page(doc, revision)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    decorate_section(section, "02 / Finance")
    add_budget_page(doc, revision)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    decorate_section(section, "03 / HR form")
    add_resume_page(doc)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=True)
    decorate_section(section, "04 / Operations")
    add_landscape_page(doc, revision)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    decorate_section(section, "05 / Governance")
    add_approval_page(doc)

    if revision > 1:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        configure_section(section)
        decorate_section(section, "06 / Appendix")
        add_appendix_page(doc)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    build_document(args.output_dir / "word-fixture-v1.docx", 1)
    build_document(args.output_dir / "word-fixture-v2.docx", 2)
    print(args.output_dir / "word-fixture-v1.docx")
    print(args.output_dir / "word-fixture-v2.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
