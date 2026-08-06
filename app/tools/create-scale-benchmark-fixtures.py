from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


def create_reference_image(path: Path) -> None:
    if path.exists():
        return
    noise = Image.effect_noise((1400, 900), 42).convert("RGB")
    overlay = Image.new("RGBA", noise.size, (27, 61, 99, 72))
    image = Image.alpha_composite(noise.convert("RGBA"), overlay)
    draw = ImageDraw.Draw(image)
    for index in range(0, 1400, 140):
        draw.rectangle((index, 0, index + 65, 900), fill=(35, 105, 155, 34))
    image.convert("RGB").save(path, "JPEG", quality=88, optimize=True)


def draw_pdf_page(c: canvas.Canvas, image: ImageReader, source_index: int, page_index: int, version: int) -> None:
    width, height = landscape(A4)
    c.setPageSize((width, height))
    c.setFillColor(colors.HexColor("#263A5A"))
    c.rect(0, height - 28, width, 28, fill=True, stroke=False)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(24, height - 19, f"Scale benchmark / PDF {source_index:02d} / Page {page_index:03d}")
    c.drawImage(image, 24, height - 250, width=270, height=174, preserveAspectRatio=True, mask="auto")

    changed = version == 2 and source_index == 1 and page_index == 11
    c.setFillColor(colors.HexColor("#1F2937"))
    c.setFont("Helvetica-Bold", 18)
    c.drawString(320, height - 78, "Cross-department operating review")
    c.setFont("Helvetica", 9)
    narrative = "Revision two updates the throughput target and escalation owner." if changed else "Baseline throughput and escalation ownership remain under review."
    c.drawString(320, height - 102, narrative)

    top = height - 280
    headers = ["Item", "Department", "Plan", "Actual", "Variance", "Status"]
    widths = [52, 128, 70, 70, 70, 92]
    x0 = 24
    row_height = 18
    c.setFont("Helvetica-Bold", 8)
    x = x0
    for header, cell_width in zip(headers, widths):
        c.setFillColor(colors.HexColor("#DCE6F1"))
        c.rect(x, top, cell_width, row_height, fill=True, stroke=True)
        c.setFillColor(colors.HexColor("#1F2937"))
        c.drawString(x + 4, top + 6, header)
        x += cell_width
    c.setFont("Helvetica", 8)
    for row in range(16):
        y = top - (row + 1) * row_height
        plan = 1000 + source_index * 70 + page_index * 9 + row * 13
        actual = plan - 18 + (37 if changed and row == 8 else (row % 5) * 7)
        values = [str(row + 1), f"Department {(row % 8) + 1}", f"{plan:,}", f"{actual:,}", f"{actual - plan:,}", "Review" if actual < plan else "On plan"]
        x = x0
        for value, cell_width in zip(values, widths):
            c.setFillColor(colors.white if row % 2 == 0 else colors.HexColor("#F5F7FA"))
            c.rect(x, y, cell_width, row_height, fill=True, stroke=True)
            c.setFillColor(colors.HexColor("#111827"))
            c.drawString(x + 4, y + 6, value)
            x += cell_width

    c.setFont("Helvetica", 7)
    c.setFillColor(colors.HexColor("#607D8B"))
    c.drawRightString(width - 24, 18, f"Synthetic load fixture - version {version}")
    c.showPage()


def create_pdf(path: Path, image_path: Path, source_index: int, page_count: int, version: int) -> None:
    c = canvas.Canvas(str(path), pagesize=landscape(A4), pageCompression=1)
    image = ImageReader(str(image_path))
    for page_index in range(1, page_count + 1):
        draw_pdf_page(c, image, source_index, page_index, version)
    c.save()


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def create_word(path: Path, image_path: Path, source_index: int, page_count: int, version: int) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)

    for page_index in range(1, page_count + 1):
        heading = document.add_paragraph()
        run = heading.add_run(f"Department evidence pack {source_index:02d} / Page {page_index:02d}")
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(38, 58, 90)

        changed = version == 2 and source_index == 1 and page_index == 5
        paragraph = document.add_paragraph(
            "Revision two changes the accountable owner and the capacity assumption for this section."
            if changed
            else "This page combines narrative, a practical table, and an embedded image for Office conversion load testing."
        )
        paragraph.paragraph_format.space_after = Pt(5)
        document.add_picture(str(image_path), width=Cm(7.2))

        table = document.add_table(rows=1, cols=6)
        table.style = "Light Shading Accent 1"
        headers = ["ID", "Deliverable", "Owner", "Plan", "Actual", "Status"]
        for column, header in enumerate(headers):
            set_cell_text(table.rows[0].cells[column], header, bold=True)
        for row_index in range(10):
            row = table.add_row()
            plan = 500 + page_index * 11 + row_index * 7
            actual = plan + (25 if changed and row_index == 4 else row_index % 4 - 2)
            values = [
                f"{page_index:02d}-{row_index + 1:02d}",
                f"Evidence item {row_index + 1}",
                f"Owner {(row_index % 5) + 1}",
                str(plan),
                str(actual),
                "Review" if actual < plan else "Complete",
            ]
            for column, value in enumerate(values):
                set_cell_text(row.cells[column], value)
        if page_index < page_count:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    parser.add_argument("--pdf-count", type=int, default=6)
    parser.add_argument("--pdf-pages", type=int, default=20)
    parser.add_argument("--word-count", type=int, default=3)
    parser.add_argument("--word-pages", type=int, default=10)
    parser.add_argument("--version", type=int, choices=(1, 2), default=1)
    parser.add_argument("--update-only", action="store_true")
    args = parser.parse_args()

    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    image_path = output / "benchmark-reference.jpg"
    create_reference_image(image_path)

    pdf_indexes = [1] if args.update_only else range(1, args.pdf_count + 1)
    word_indexes = [1] if args.update_only else range(1, args.word_count + 1)
    for source_index in pdf_indexes:
        create_pdf(output / f"large-pdf-{source_index:02d}.pdf", image_path, source_index, args.pdf_pages, args.version)
    for source_index in word_indexes:
        create_word(output / f"large-word-{source_index:02d}.docx", image_path, source_index, args.word_pages, args.version)

    image_path.unlink(missing_ok=True)
    print(f"created scale fixtures in {output}")


if __name__ == "__main__":
    main()
